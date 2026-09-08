"""
적용 내역 리포트 테스트.

사용자가 이 리포트로 확인하려는 것은 하나다 — "에이전트가 정말 일관되게
바꿨는가". 그러려면 표시가 **정확한 자리에만** 찍혀야 한다. 'set'이
'settings' 안에서 잡히거나, 긴 표현이 짧은 표현에 부서지면 리포트가
오히려 신뢰를 깎는다.

    python tests/test_effect_report.py
"""
from __future__ import annotations

import io
import sys
import tempfile
from pathlib import Path

ROOT = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(ROOT))

from docx import Document

from services import effect_report as er

failures = []


def check(name, cond, detail=""):
    print(f"  {'OK  ' if cond else 'FAIL'} {name}" + (f"  {detail}" if not cond else ""))
    if not cond:
        failures.append(name)


print("[1] 표시 위치 — 단어 경계")
sp = er.split_spans("Open the settings and set the value.", [("set", "글로서리")])
marked = [t for t, k in sp if k]
check("settings 안의 set은 잡지 않는다", marked == ["set"], str(marked))

print("[2] 긴 표현이 먼저")
sp = er.split_spans(
    "Upload the login record file to the file browser.",
    [("file", "글로서리"), ("login record file", "글로서리")],
)
marked = [t for t, k in sp if k]
check("긴 것부터 잡는다", "login record file" in marked, str(marked))
check("남은 file도 잡는다", marked.count("file") == 1, str(marked))
check("겹치지 않는다", len(marked) == 2, str(marked))

print("[3] 원문이 손실되지 않는다")
text = "Click **Search File** then Save."
sp = er.split_spans(text, [("Search File", "UI 매핑"), ("Save", "UI 매핑")])
check("조각을 이으면 원문", "".join(t for t, _ in sp) == text,
      "".join(t for t, _ in sp))

print("[4] 대소문자 무시, 원문 표기 유지")
sp = er.split_spans("analysis and Analysis", [("Analysis", "글로서리")])
marked = [t for t, k in sp if k]
check("둘 다 잡는다", len(marked) == 2, str(marked))
check("원문 표기 그대로", marked == ["analysis", "Analysis"], str(marked))

print("[5] 출처가 조각에 실린다")
sp = er.split_spans("Save and analysis",
                    [("Save", "UI 매핑"), ("analysis", "글로서리")])
kinds = {t: k for t, k in sp if k}
check("UI 매핑 구분", kinds.get("Save") == "UI 매핑", str(kinds))
check("글로서리 구분", kinds.get("analysis") == "글로서리", str(kinds))

print("[6] 산출물 읽기 — 마크다운")
tmp = Path(tempfile.mkdtemp(prefix="effect_"))
md = tmp / "out.mdx"
md.write_text("# Title\n\nFirst line.\n\nSecond line.\n", encoding="utf-8")
paras = er.read_output_paragraphs(str(md))
check("빈 줄 제외", paras == ["# Title", "First line.", "Second line."], str(paras))

print("[7] 리포트 생성")
applied = [
    {"KO": "분석", "EN": "analysis", "출처": "글로서리", "적용": 2, "예문": ""},
    {"KO": "저장", "EN": "Save", "출처": "UI 매핑", "적용": 1, "예문": ""},
]
body = tmp / "body.mdx"
body.write_text("Run the analysis.\n\nClick Save to finish.\n\nNothing here.\n",
                encoding="utf-8")
data = er.build(applied, str(body), "runAnalysis.mdx", "Sparrow")
check("docx 바이트", data[:2] == b"PK" and len(data) > 5000, str(len(data)))

doc = Document(io.BytesIO(data))
check("표 1개", len(doc.tables) == 1, str(len(doc.tables)))
check("표에 머리행 + 2건", len(doc.tables[0].rows) == 3,
      str(len(doc.tables[0].rows)))

hl = [(r.text, str(r.font.highlight_color))
      for p in doc.paragraphs for r in p.runs if r.font.highlight_color]
texts = [t for t, _ in hl]
check("본문 표현이 표시된다", "analysis" in texts and "Save" in texts, str(texts))
check("Glossary는 노란색",
      any(t == "analysis" and "YELLOW" in c for t, c in hl), str(hl))
check("UI 매핑은 초록색",
      any(t == "Save" and "BRIGHT_GREEN" in c for t, c in hl), str(hl))

_all = "\n".join(p.text for p in doc.paragraphs)
check("적용 없는 문단은 빠진다", "Nothing here." not in _all)
check("생략 안내", "3개 문단 중" in _all or "문단만 수록" in _all, _all[-160:])
check("요약 문장", "2건이 본문" in _all, _all[:400])

print("[8] 적용 내역이 비어도 죽지 않는다")
data0 = er.build([], str(body), "x.mdx", None)
check("빈 입력도 문서 생성", data0[:2] == b"PK", str(len(data0)))

import shutil  # noqa: E402
shutil.rmtree(tmp, ignore_errors=True)

print()
if failures:
    print(f"FAILED {len(failures)}건: {failures}")
    raise SystemExit(1)
print("ALL PASS")
